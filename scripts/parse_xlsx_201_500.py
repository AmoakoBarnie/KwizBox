"""Parse Ghana_STEM_Trivia_Questions_201-500.xlsx -> questions_201_500.json
Compatible schema with the existing questions.json (ids 1-200):
  id, class_level, subject, topic, strand, difficulty, question, options[4],
  answer (letter), answer_index (0..3), explanation
"""
import json
import sys
import openpyxl

XLSX = sys.argv[1] if len(sys.argv) > 1 else "Ghana_STEM_Trivia_Questions_201-500.xlsx"
OUT = sys.argv[2] if len(sys.argv) > 2 else "questions_201_500.json"

CLASS_MAP = {"B4": "B4", "B5": "B5", "B6": "B6", "B7": "B7", "B8": "B8", "B9": "B9"}
SUBJECT_MAP = {"Mathematics": "Mathematics", "Science": "Science", "Computing": "Computing"}
DIFF_MAP = {"Easy": "Easy", "Medium": "Medium", "Hard": "Hard"}
LETTER_IDX = {"A": 0, "B": 1, "C": 2, "D": 3}


def norm(s):
    return (s or "").strip()


def main():
    wb = openpyxl.load_workbook(XLSX, data_only=True)
    ws = wb.worksheets[0]

    rows = []
    problems = []
    for r in range(2, ws.max_row + 1):
        qnum = ws.cell(row=r, column=1).value
        if qnum is None:
            continue  # spacer row
        qnum = int(qnum)
        cls = norm(ws.cell(row=r, column=2).value)
        subj = norm(ws.cell(row=r, column=3).value)
        topic = norm(ws.cell(row=r, column=4).value)
        diff = norm(ws.cell(row=r, column=5).value)
        question = norm(ws.cell(row=r, column=6).value)
        opts = [norm(ws.cell(row=r, column=c).value) for c in range(7, 11)]
        answer_letter = norm(ws.cell(row=r, column=11).value)
        explanation = norm(ws.cell(row=r, column=12).value)

        # Validate
        if cls not in CLASS_MAP:
            problems.append((qnum, f"bad class '{cls}'"))
        if subj not in SUBJECT_MAP:
            problems.append((qnum, f"bad subject '{subj}'"))
        if diff not in DIFF_MAP:
            problems.append((qnum, f"bad difficulty '{diff}'"))
        if answer_letter not in LETTER_IDX:
            problems.append((qnum, f"bad answer letter '{answer_letter}'"))
        if not question:
            problems.append((qnum, "empty question"))
        if not explanation:
            problems.append((qnum, "empty explanation"))

        ai = LETTER_IDX.get(answer_letter)
        # Note: source has some duplicate option values (e.g. B and C both "75").
        # We trust the Answer LETTER as source of truth; just record it.
        strand = topic.split("/")[0].strip() if "/" in topic else topic

        # Skip rows that are missing essential data (e.g. empty option text) — we
        # can't faithfully invent answer choices, so we exclude them and report.
        if any(o == "" for o in opts) or ai is None or not question or not explanation:
            problems.append((qnum, "skipped: missing option/answer/question/explanation"))
            continue

        rows.append({
            "id": qnum,
            "class_level": CLASS_MAP.get(cls, cls),
            "subject": SUBJECT_MAP.get(subj, subj),
            "topic": topic,
            "strand": strand,
            "difficulty": DIFF_MAP.get(diff, diff),
            "question": question,
            "options": opts,
            "answer": answer_letter,
            "answer_index": ai,
            "explanation": explanation,
        })

    # Check id continuity 201..500
    ids = [x["id"] for x in rows]
    expected = list(range(201, 501))
    if ids != expected:
        missing = set(expected) - set(ids)
        problems.append((0, f"id gap, missing {sorted(missing)[:10]}..." if missing else "ids out of order"))

    print(f"Parsed {len(rows)} questions (Q# {ids[0]}..{ids[-1]})")
    if problems:
        print(f"!! {len(problems)} PROBLEMS:")
        for p in problems[:40]:
            print("   ", p)
    else:
        print("No validation problems.")

    with open(OUT, "w", encoding="utf-8") as f:
        json.dump(rows, f, ensure_ascii=False, indent=2)
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    main()
